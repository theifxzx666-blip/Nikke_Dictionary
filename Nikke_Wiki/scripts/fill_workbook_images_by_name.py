"""Insert missing workbook images using exact character-name matches."""

import argparse
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "nikke_character_dictionary_enhanced.json"
CANONICAL = ROOT / "data" / "character_repository" / "canonical" / "characters.json"
# 历史重建管线产物，2026-10-21 起归档到 data/archive/legacy_exports_20261021/
LEGACY_EXPORT_DIR = ROOT / "data" / "archive" / "legacy_exports_20261021"
INPUT = LEGACY_EXPORT_DIR / "nikke_character_dictionary_rebuilt_filled.xlsx"
OUTPUT = LEGACY_EXPORT_DIR / "nikke_character_dictionary_rebuilt_filled_with_images.xlsx"
CACHE = ROOT / "assets" / ".nikke_xlsx_images_repair"


def src_map(rows):
    out = {}
    for row in rows:
        for key in (row.get("name"), row.get("cnName"), row.get("gamekeeCnName")):
            if key:
                out[str(key).strip()] = row
    return out


def image_cache(source: str, key: str, max_size: tuple[int, int]) -> Path | None:
    if not source or not Path(source).is_file():
        return None
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / f"{key}.png"
    if not out.exists():
        try:
            with Image.open(source) as image:
                image = image.convert("RGBA")
                image.thumbnail(max_size, Image.Resampling.LANCZOS)
                canvas = Image.new("RGBA", max_size, (255, 255, 255, 0))
                canvas.alpha_composite(image, ((max_size[0] - image.width) // 2, (max_size[1] - image.height) // 2))
                canvas.save(out)
        except (OSError, ValueError):
            return None
    return out


def existing_anchors(ws):
    result = set()
    for image in ws._images:
        anchor = image.anchor
        if hasattr(anchor, "_from"):
            result.add((anchor._from.row + 1, anchor._from.col + 1))
    return result


def add_if_missing(ws, row_num, col_num, source, key, size, display, stats):
    if (row_num, col_num) in stats["anchors"].setdefault(ws.title, set()):
        return
    path = image_cache(source, key, size)
    if not path:
        stats["missing"] += 1
        return
    image = XlsxImage(str(path))
    image.width, image.height = display
    ws.add_image(image, ws.cell(row=row_num, column=col_num).coordinate)
    stats["inserted"] += 1
    stats["anchors"][ws.title].add((row_num, col_num))


def main() -> None:
    parser = argparse.ArgumentParser(description="Fill missing character images in rebuilt workbook")
    parser.add_argument("--input", type=Path, default=INPUT)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    rows = json.loads(SOURCE.read_text(encoding="utf-8"))
    by_name = src_map(rows)
    canonical = json.loads(CANONICAL.read_text(encoding="utf-8")) if CANONICAL.exists() else {"characters": []}
    canonical_by_slug = {item["identity"]["dictionarySlug"]: item for item in canonical.get("characters", [])}
    workbook = load_workbook(args.input)
    stats = {"inserted": 0, "missing": 0, "anchors": {ws.title: existing_anchors(ws) for ws in workbook.worksheets}, "sheetInserted": {}}

    def process(sheet_name, name_labels, specs, replace_fields=()):
        ws = workbook[sheet_name]
        headers = [cell.value for cell in ws[1]]
        index = {str(label): i + 1 for i, label in enumerate(headers)}
        if replace_fields:
            replace_cols = {index[label] for label in replace_fields if label in index}
            ws._images = [image for image in ws._images if not (hasattr(image.anchor, "_from") and image.anchor._from.col + 1 in replace_cols)]
            stats["anchors"][sheet_name] = existing_anchors(ws)
        before = stats["inserted"]
        for row_num in range(2, ws.max_row + 1):
            source = None
            for label in name_labels:
                col = index.get(label)
                if col and ws.cell(row=row_num, column=col).value:
                    source = by_name.get(str(ws.cell(row=row_num, column=col).value).strip())
                    if source:
                        break
            if not source:
                continue
            for label, field, size, display in specs:
                col = index.get(label)
                if col:
                    path = source.get(field, "")
                    if field == "localPortraitFile":
                        gg = canonical_by_slug.get(source.get("url", ""), {}).get("assets", {}).get("nikkeGgPortrait", {})
                        if gg.get("status") == "available":
                            path = str(ROOT / "data" / "character_repository" / gg.get("file", "").replace("assets/characters/", "assets/characters/"))
                    add_if_missing(ws, row_num, col, path, f"{sheet_name}_{row_num}_{field}", size, display, stats)
        stats["sheetInserted"][sheet_name] = stats["inserted"] - before

    process("角色字典", ["英文名", "中文名"], [("头像", "localIconFile", (128, 128), (64, 64))])
    process("全汉化", ["英文名", "角色名"], [("头像", "localIconFile", (128, 128), (64, 64)), ("立绘", "localPortraitFile", (128, 128), (72, 100)), ("属性图标", "gamekeeLocalElementIconFile", (64, 64), (34, 34))], replace_fields=("立绘",))
    process("国服上线状态", ["英文名", "角色名"], [("头像", "localIconFile", (128, 128), (64, 64))])
    process("GameKee素材", ["英文名", "角色名"], [
        ("头像", "localIconFile", (128, 128), (54, 70)),
        ("配队卡", "gamekeeLocalTeamCardFile", (128, 128), (54, 70)),
        ("立绘", "localPortraitFile", (128, 128), (72, 100)),
        ("时装立绘", "gamekeeLocalCostumePortraitFile", (128, 128), (72, 100)),
        ("阶段图标", "gamekeeLocalStageIconFile", (64, 64), (34, 34)),
        ("职业图标", "gamekeeLocalClassIconFile", (64, 64), (34, 34)),
        ("企业图标", "gamekeeLocalCompanyIconFile", (64, 64), (34, 34)),
        ("武器图标", "gamekeeLocalWeaponIconFile", (64, 64), (34, 34)),
        ("属性图标", "gamekeeLocalElementIconFile", (64, 64), (34, 34)),
    ], replace_fields=("立绘",))
    process("皮肤清单", ["英文名", "角色名"], [("皮肤立绘", "localPortraitFile", (128, 128), (72, 100))])
    process("NPC与剧情角色", ["英文名", "角色名"], [("角色卡", "localPortraitFile", (128, 128), (54, 100))])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    report = {"input": str(args.input), "output": str(args.output), "insertedImages": stats["inserted"], "missingImages": stats["missing"], "sheetInserted": stats["sheetInserted"]}
    report_path = ROOT / "data" / "character_repository" / "reports" / "workbook_image_fill_latest.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
