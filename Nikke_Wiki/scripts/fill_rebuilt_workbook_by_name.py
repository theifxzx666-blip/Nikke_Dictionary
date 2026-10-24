"""Fill blank cells in the cleaned rebuilt workbook using exact character-name matches."""

import argparse
import copy
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "nikke_character_dictionary_enhanced.json"
# 历史重建管线产物，2026-10-21 起归档到 data/archive/legacy_exports_20261021/
LEGACY_EXPORT_DIR = ROOT / "data" / "archive" / "legacy_exports_20261021"
DEFAULT_INPUT = LEGACY_EXPORT_DIR / "nikke_character_dictionary_rebuilt.xlsx"
DEFAULT_OUTPUT = LEGACY_EXPORT_DIR / "nikke_character_dictionary_rebuilt_filled.xlsx"


def text(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "|".join(str(item) for item in value)
    return str(value)


def skill_text(skill: dict) -> str:
    if not skill:
        return ""
    parts = [skill.get("name", ""), skill.get("type", ""), skill.get("cooldown", ""), skill.get("effect", "")]
    return "\n".join(str(value) for value in parts if value)


def source_map(rows):
    result = {}
    for row in rows:
        for key in (row.get("name"), row.get("cnName"), row.get("gamekeeCnName")):
            if key:
                result[str(key).strip()] = row
    return result


def fill_cell(cell, value, stats):
    if value in (None, "", []):
        return
    if cell.value is None or str(cell.value).strip() == "":
        cell.value = value
        stats["filledCells"] += 1


def row_source(row, headers, mapping, by_name):
    indices = {str(label): i for i, label in enumerate(headers)}
    for label in ("英文名", "角色名", "角色名称", "角色名", "GameKee名称"):
        idx = indices.get(label)
        if idx is not None and row[idx].value:
            source = by_name.get(str(row[idx].value).strip())
            if source:
                return source
    return None


def main() -> None:
    parser = argparse.ArgumentParser(description="Fill cleaned NIKKE workbook cells by exact character names")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    rows = json.loads(SOURCE.read_text(encoding="utf-8"))
    by_name = source_map(rows)
    workbook = load_workbook(args.input)
    stats = {"sourceCharacters": len(rows), "filledCells": 0, "unmatchedRows": [], "sheetFilled": {}}

    sheet_maps = {
        "角色字典": {"英文名": "name", "中文名": "cnName", "别名": "aliases", "国服是否上线": "cnReleased", "爆裂": "burst", "属性": "element", "职业": "class", "武器": "weapon", "企业": "manufacturer", "部队": "squad", "稀有度": "rarity", "Nikke.gg Slug": "url", "Nikke.gg头像": "dotggIconUrl", "GameKee头像": "gamekeeIconUrl"},
        "全汉化": {"角色名": "cnName", "常用别名": "aliases", "国服是否上线": "cnReleased", "定位": "gamekeeClassCn", "武器": "gamekeeWeaponCn", "企业": "gamekeeCompanyCn", "部队": "gamekeeSquadCn", "稀有度": "gamekeeRarityCn", "英文名": "name", "GameKee名称": "gamekeeCnName", "Wiki校验": "wikiNameStatus", "Nikke.gg头像": "dotggIconUrl", "GameKee头像": "gamekeeIconUrl", "GameKee立绘": "gamekeePortraitUrl"},
        "国服上线状态": {"角色名": "cnName", "英文名": "name", "国服是否上线": "cnReleased", "判断依据": "releaseBasis", "Wiki状态": "wikiNameStatus", "GameKee条目ID": "gamekeeEntryId", "GameKee名称": "gamekeeCnName", "别名": "aliases"},
        "GameKee素材": {"角色名": "cnName", "英文名": "name", "阶段": "gamekeeStageCn", "职业": "gamekeeClassCn", "企业": "gamekeeCompanyCn", "武器": "gamekeeWeaponCn", "属性": "gamekeeElementCn", "配队卡URL": "gamekeeTeamCardUrl", "立绘URL": "gamekeePortraitUrl", "时装立绘URL": "gamekeeCostumePortraitUrl"},
    }
    for sheet_name, field_map in sheet_maps.items():
        ws = workbook[sheet_name]
        headers = [cell.value for cell in ws[1]]
        sheet_stats = 0
        for row in ws.iter_rows(min_row=2):
            source = row_source(row, headers, field_map, by_name)
            if not source:
                stats["unmatchedRows"].append({"sheet": sheet_name, "row": row[0].row})
                continue
            for cell, label in zip(row, headers):
                field = field_map.get(str(label))
                if field:
                    before = stats["filledCells"]
                    fill_cell(cell, text(source.get(field, "")), stats)
                    sheet_stats += stats["filledCells"] - before
        stats["sheetFilled"][sheet_name] = sheet_stats

    ws = workbook["技能组"]
    headers = [cell.value for cell in ws[1]]
    labels = {str(label): i for i, label in enumerate(headers)}
    for row in ws.iter_rows(min_row=2):
        source = row_source(row, headers, {}, by_name)
        if not source:
            stats["unmatchedRows"].append({"sheet": "技能组", "row": row[0].row})
            continue
        values = {"角色名": source.get("cnName"), "英文名": source.get("name"), "所属部队": source.get("gamekeeSquadCn"), "爆裂阶段": source.get("gamekeeStageCn"), "枪械类型": source.get("gamekeeWeaponCn"), "普通攻击": json.dumps(source.get("normalAttack", {}), ensure_ascii=False), "技能1": skill_text(source.get("skills", {}).get("skill1", {})), "技能2": skill_text(source.get("skills", {}).get("skill2", {})), "爆裂技能": skill_text(source.get("skills", {}).get("burstSkill", {})), "资料来源": source.get("source", {}).get("url", "")}
        for label, value in values.items():
            if label in labels:
                fill_cell(row[labels[label]], text(value), stats)

    ws = workbook["珍藏品"]
    headers = [cell.value for cell in ws[1]]
    labels = {str(label): i for i, label in enumerate(headers)}
    for row in ws.iter_rows(min_row=2):
        source = row_source(row, headers, {}, by_name)
        if not source:
            continue
        favorite = source.get("favoriteItem", {})
        values = {"角色名": source.get("cnName"), "英文名": source.get("name"), "珍藏品状态": favorite.get("status"), "珍藏品名称": favorite.get("itemName"), "国服状态": favorite.get("cnReleased"), "说明": favorite.get("note"), "珍藏品强化技能": json.dumps(favorite.get("skills", {}), ensure_ascii=False), "资料来源": favorite.get("sourceUrl", "")}
        for label, value in values.items():
            if label in labels:
                fill_cell(row[labels[label]], text(value), stats)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    report_path = ROOT / "data" / "character_repository" / "reports" / "workbook_fill_latest.json"
    report_path.write_text(json.dumps({"input": str(args.input), "output": str(args.output), **stats}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"input": str(args.input), "output": str(args.output), "filledCells": stats["filledCells"], "sheetFilled": stats["sheetFilled"], "unmatchedRows": len(stats["unmatchedRows"])}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
