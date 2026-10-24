import base64
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = PROJECT_ROOT / "data"
ASSET_DIR = PROJECT_ROOT / "assets"
EXPORT_DIR = PROJECT_ROOT / "正式产物"
ENHANCED_DICTIONARY = DATA_DIR / "nikke_character_dictionary_enhanced.json"
CN_RELEASE_STATUS = DATA_DIR / "cn_release_status.json"
NON_PLAYABLE_CARDS = DATA_DIR / "non_playable_character_cards.json"
ICON_DIR = ASSET_DIR / "nikke_character_icons"
CODE_ICON_DIR = ASSET_DIR / "nikke_code_icons"
GAMEKEE_CODE_ICON_DIR = ASSET_DIR / "nikke_code_icons_gamekee"
HTML_OUT = EXPORT_DIR / "nikke_character_dictionary_with_images.html"
XLSX_OUT = EXPORT_DIR / "nikke_character_dictionary_with_images.xlsx"
XLSX_FALLBACK_OUT = EXPORT_DIR / "nikke_character_dictionary_with_images_updated.xlsx"
TMP_DIR = ASSET_DIR / ".nikke_xlsx_icons"


RAW_HEADERS = [
    ("avatar", "头像"),
    ("name", "英文名"),
    ("cnName", "中文名"),
    ("aliases", "别名"),
    ("cnReleased", "国服是否上线"),
    ("burst", "爆裂"),
    ("element", "属性"),
    ("class", "职业"),
    ("weapon", "武器"),
    ("manufacturer", "企业"),
    ("squad", "部队"),
    ("rarity", "稀有度"),
    ("url", "Nikke.gg Slug"),
    ("dotggIconUrl", "Nikke.gg头像"),
    ("gamekeeIconUrl", "GameKee头像"),
]

CN_HEADERS = [
    ("avatar", "头像"),
    ("portrait", "立绘"),
    ("cnName", "角色名"),
    ("aliases", "常用别名"),
    ("cnReleased", "国服是否上线"),
    ("burstCn", "爆裂"),
    ("elementIcon", "属性图标"),
    ("elementCn", "属性代码"),
    ("classCn", "定位"),
    ("weaponCn", "武器"),
    ("manufacturerCn", "企业"),
    ("squad", "部队"),
    ("rarityCn", "稀有度"),
    ("name", "英文名"),
    ("gamekeeCnName", "GameKee名称"),
    ("wikiNameStatusCn", "Wiki校验"),
    ("dotggIconUrl", "Nikke.gg头像"),
    ("gamekeeIconUrl", "GameKee头像"),
    ("gamekeePortraitUrl", "GameKee立绘"),
]

STATUS_HEADERS = [
    ("avatar", "头像"),
    ("cnName", "角色名"),
    ("name", "英文名"),
    ("cnReleased", "国服是否上线"),
    ("releaseBasis", "判断依据"),
    ("wikiNameStatusCn", "Wiki状态"),
    ("gamekeeEntryId", "GameKee条目ID"),
    ("gamekeeCnName", "GameKee名称"),
    ("aliases", "别名"),
]

ASSET_HEADERS = [
    ("avatar", "头像"),
    ("teamCard", "配队卡"),
    ("portrait", "立绘"),
    ("costumePortrait", "时装立绘"),
    ("cnName", "角色名"),
    ("name", "英文名"),
    ("gamekeeStageCn", "阶段"),
    ("gamekeeStageIcon", "阶段图标"),
    ("gamekeeClassCn", "职业"),
    ("gamekeeClassIcon", "职业图标"),
    ("gamekeeCompanyCn", "企业"),
    ("gamekeeCompanyIcon", "企业图标"),
    ("gamekeeWeaponCn", "武器"),
    ("gamekeeWeaponIcon", "武器图标"),
    ("gamekeeElementCn", "属性"),
    ("gamekeeElementIcon", "属性图标"),
    ("gamekeeTeamCardUrl", "配队卡URL"),
    ("gamekeePortraitUrl", "立绘URL"),
    ("gamekeeCostumePortraitUrl", "时装立绘URL"),
]

SKILL_HEADERS = [
    ("cnName", "角色名"),
    ("name", "英文名"),
    ("gamekeeSquadCn", "所属部队"),
    ("gamekeeStageCn", "爆裂阶段"),
    ("gamekeeWeaponCn", "枪械类型"),
    ("normalAttack", "普通攻击"),
    ("skill1", "技能1"),
    ("skill2", "技能2"),
    ("burstSkill", "爆裂技能"),
    ("sourceUrl", "资料来源"),
]

FAVORITE_ITEM_HEADERS = [
    ("cnName", "角色名"),
    ("name", "英文名"),
    ("status", "珍藏品状态"),
    ("itemName", "珍藏品名称"),
    ("cnReleased", "国服状态"),
    ("note", "说明"),
    ("skills", "珍藏品强化技能"),
    ("sourceUrl", "资料来源"),
]

COSTUME_HEADERS = [
    ("parentCnName", "\u89d2\u8272\u540d"),
    ("parentName", "\u82f1\u6587\u540d"),
    ("portrait", "\u76ae\u80a4\u7acb\u7ed8"),
    ("name", "\u76ae\u80a4\u540d\u79f0"),
    ("status", "\u7d20\u6750\u72b6\u6001"),
    ("contentId", "GameKee\u5185\u5bb9ID"),
    ("sourceUrl", "\u6765\u6e90\u9875"),
    ("portraitUrl", "\u7acb\u7ed8URL"),
]

NON_PLAYABLE_HEADERS = [("portrait", "角色卡"), ("name", "角色名"), ("aliases", "别名"), ("company", "企业"), ("squad", "部队"), ("stage", "爆裂阶段"), ("element", "属性"), ("class", "职业"), ("weapon", "武器"), ("skills", "技能组"), ("sourceUrl", "来源")]

OFFICIAL_ASSET_HEADERS = [
    ("portrait", "官方头像"),
    ("cnName", "角色名"),
    ("name", "英文名"),
    ("code", "素材代号"),
    ("status", "收录状态"),
    ("portraitSize", "头像规格"),
    ("fullbodySize", "立绘规格"),
    ("hdSize", "HD规格"),
    ("fullbodyPath", "立绘路径"),
    ("hdPath", "HD路径"),
    ("source", "来源"),
]

ELEMENT_CN = {
    "Fire": "燃烧",
    "Water": "水冷",
    "Wind": "风压",
    "Electric": "电击",
    "Iron": "铁甲",
}

CLASS_CN = {
    "Attacker": "火力型",
    "Defender": "防御型",
    "Supporter": "辅助型",
}

MANUFACTURER_CN = {
    "Elysion": "极乐净土",
    "Missilis": "米西利斯",
    "Tetra": "泰特拉",
    "Pilgrim": "朝圣者",
    "Abnormal": "反常者",
}

WEAPON_CN = {
    "AR": "步枪",
    "MG": "机枪",
    "SMG": "冲锋枪",
    "SG": "霰弹枪",
    "SR": "狙击枪",
    "RL": "发射器",
}

BURST_CN = {
    "1": "爆裂Ⅰ",
    "2": "爆裂Ⅱ",
    "3": "爆裂Ⅲ",
    "p": "全爆裂",
}


def normalize_rows(rows: list[dict]) -> list[dict]:
    release_data = {}
    if CN_RELEASE_STATUS.exists():
        raw = json.loads(CN_RELEASE_STATUS.read_text(encoding="utf-8"))
        release_data = {k: v for k, v in raw.items() if not k.startswith("_")}
    for row in rows:
        info = release_data.get(row.get("name", ""), {})
        row["cnReleased"] = info.get("status", "待确认")
        basis = info.get("basis", "")
        date = info.get("date", "")
        note = info.get("note", "")
        parts = []
        if basis:
            parts.append(basis)
        if date:
            parts.append(f"上线日期：{date}")
        if note:
            parts.append(note)
        row["releaseBasis"] = "；".join(parts) if parts else "未录入国服上线状态"
    return rows


def icon_path(row: dict) -> Path | None:
    local_text = row.get("localIconFile", "")
    if not local_text:
        return None
    local = Path(local_text)
    if not local.is_file():
        return None
    try:
        with Image.open(local) as image:
            return local if image.size == (128, 128) else None
    except (OSError, ValueError):
        return None


def code_icon_path(row: dict) -> Path | None:
    gamekee_path = GAMEKEE_CODE_ICON_DIR / f"{row.get('element', '')}.png"
    if gamekee_path.exists():
        TMP_DIR.mkdir(exist_ok=True)
        out = TMP_DIR / f"code_{row.get('element', '')}.png"
        if not out.exists():
            Image.open(gamekee_path).convert("RGBA").save(out)
        return out
    path = CODE_ICON_DIR / f"{row.get('element', '')}.png"
    return path if path.exists() else None


def data_uri(path: Path) -> str:
    mime = "image/webp" if path.suffix.lower() == ".webp" else "image/png"
    data = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{data}"


def export_html(rows: list[dict]) -> None:
    html = [
        "<!doctype html>",
        "<html lang='zh-CN'>",
        "<head><meta charset='utf-8'><title>NIKKE 角色字典</title>",
        "<style>",
        "body{font-family:Arial,'Microsoft YaHei',sans-serif;margin:0;background:#eef2f3;color:#20272d}",
        "header{padding:24px 32px 18px;background:#fff;border-bottom:4px solid #d92843}",
        "h1{margin:0;font-size:24px;font-weight:700}.muted{color:#65717c;margin:6px 0 0}",
        ".table-wrap{padding:20px 32px 32px;overflow:auto}",
        "table{border-collapse:separate;border-spacing:0;width:100%;background:#fff;font-size:13px;box-shadow:0 1px 3px #20272d1a}",
        "th,td{border-bottom:1px solid #dfe5e7;padding:8px 10px;text-align:left;vertical-align:middle;white-space:nowrap}",
        "th{position:sticky;top:0;background:#25333a;color:#fff;z-index:1;font-weight:600}",
        "tbody tr:nth-child(even){background:#f7f9f9}tbody tr:hover{background:#fff4f5}",
        ".avatar{width:64px;height:64px;object-fit:cover;border-radius:4px;background:#dde3e5;display:block}",
        "</style></head><body>",
        "<header><h1>NIKKE 角色字典</h1>",
        f"<p class='muted'>共 {len(rows)} 个角色。头像为 Nikke.gg 标准头像，中文资料按 GameKee 图鉴校验。</p></header>",
        "<main class='table-wrap'><table><thead><tr>",
    ]
    for _, label in RAW_HEADERS:
        html.append(f"<th>{label}</th>")
    html.append("</tr></thead><tbody>")
    for row in rows:
        html.append("<tr>")
        local_icon = icon_path(row)
        img_src = data_uri(local_icon) if local_icon else (
            row.get("dotggIconUrl", "") if row.get("avatarStatus") in {"available", "available_derived"} else ""
        )
        for key, _ in RAW_HEADERS:
            if key == "avatar":
                avatar = f"<img class='avatar' src='{img_src}' alt='{row.get('name', '')}'>" if img_src else "<span class='muted'>待补头像</span>"
                html.append(f"<td>{avatar}</td>")
            else:
                html.append(f"<td>{row.get(key, '') or ''}</td>")
        html.append("</tr>")
    html.append("</tbody></table></main></body></html>")
    HTML_OUT.write_text("\n".join(html), encoding="utf-8")


def png_for_xlsx(row: dict) -> Path | None:
    src = icon_path(row)
    if not src:
        return None
    TMP_DIR.mkdir(exist_ok=True)
    if row.get("avatarSource") in {"Nikke.gg", "Nikke.gg_derived_from_full_art"}:
        # dotgg 标准头像：独立 128x128 方形缓存（避免与旧 54x70 缓存冲突）
        out = TMP_DIR / f"{row.get('url', 'unknown')}_128.png"
    else:
        # GameKee 卡片图等：沿用原 {url}.png 缓存
        out = TMP_DIR / f"{row.get('url', 'unknown')}.png"
    if out.exists():
        return out
    img = Image.open(src).convert("RGBA")
    if row.get("avatarSource") in {"Nikke.gg", "Nikke.gg_derived_from_full_art"}:
        # dotgg 标准头像：保持 128x128 方形（不拉伸）
        img.thumbnail((128, 128), Image.Resampling.LANCZOS)
        canvas = Image.new("RGBA", (128, 128), (255, 255, 255, 0))
        canvas.alpha_composite(img, ((128 - img.width) // 2, (128 - img.height) // 2))
    else:
        # GameKee 卡片图等：保持原 54x70 处理
        img.thumbnail((54, 70), Image.Resampling.LANCZOS)
        canvas = Image.new("RGBA", (54, 70), (255, 255, 255, 0))
        canvas.alpha_composite(img, ((54 - img.width) // 2, (70 - img.height) // 2))
    canvas.save(out)
    return out


def image_file_for_xlsx(path_text: str, cache_name: str, size: int = 64) -> Path | None:
    if not path_text:
        return None
    src = Path(path_text)
    if not src.exists():
        return None
    TMP_DIR.mkdir(exist_ok=True)
    out = TMP_DIR / f"{cache_name}.png"
    if out.exists():
        return out
    img = Image.open(src).convert("RGBA")
    img.thumbnail((size, size), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (size, size), (255, 255, 255, 0))
    canvas.alpha_composite(img, ((size - img.width) // 2, (size - img.height) // 2))
    canvas.save(out)
    return out


def add_image_from_path(ws, path_text: str, cell_ref: str, cache_name: str, width: int, height: int, size: int = 96) -> None:
    png = image_file_for_xlsx(path_text, cache_name, size)
    if not png:
        return
    img = XlsxImage(str(png))
    img.width = width
    img.height = height
    ws.add_image(img, cell_ref)


def add_headers(ws, headers):
    for col, (_, label) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25333A")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def finish_sheet(ws) -> None:
    border = Border(bottom=Side(style="hair", color="DDE3E5"))
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F6F8F8")


def add_avatar(ws, row: dict, cell_ref: str) -> None:
    png = png_for_xlsx(row)
    if png:
        img = XlsxImage(str(png))
        if row.get("avatarSource") == "Nikke.gg":
            # dotgg 标准头像：128x128 方形显示
            img.width = 128
            img.height = 128
        else:
            # GameKee 卡片图等：保持原 48x62 显示
            img.width = 48
            img.height = 62
        ws.add_image(img, cell_ref)
    else:
        ws[cell_ref] = "缺图"


def add_portrait(ws, row: dict, cell_ref: str) -> None:
    path_text = row.get("localPortraitFile", "")
    if not path_text:
        return
    path = Path(path_text)
    if not path.exists():
        return
    TMP_DIR.mkdir(exist_ok=True)
    out = TMP_DIR / f"portrait_{row.get('url', 'unknown')}.png"
    if not out.exists():
        img = Image.open(path).convert("RGBA")
        img.thumbnail((72, 72), Image.Resampling.LANCZOS)
        canvas = Image.new("RGBA", (72, 72), (255, 255, 255, 0))
        canvas.alpha_composite(img, ((72 - img.width) // 2, (72 - img.height) // 2))
        canvas.save(out)
    ximg = XlsxImage(str(out))
    ximg.width = 58
    ximg.height = 58
    ws.add_image(ximg, cell_ref)


def chinese_value(row: dict, key: str) -> str:
    if key == "burstCn":
        return BURST_CN.get(str(row.get("burst", "")), row.get("burst", ""))
    if key == "elementCn":
        return ELEMENT_CN.get(row.get("element", ""), row.get("element", ""))
    if key == "classCn":
        return CLASS_CN.get(row.get("class", ""), row.get("class", ""))
    if key == "weaponCn":
        return WEAPON_CN.get(row.get("weapon", ""), row.get("weapon", ""))
    if key == "manufacturerCn":
        return MANUFACTURER_CN.get(row.get("manufacturer", ""), row.get("manufacturer", ""))
    if key == "rarityCn":
        return row.get("rarity", "")
    if key == "wikiNameStatusCn":
        status = row.get("wikiNameStatus", "")
        if status == "matched":
            return "已匹配"
        if status == "gamekee_only":
            return "仅GameKee"
        return "待校验"
    if key == "dotggIconUrl":
        return row.get("dotggIconUrl") or row.get("iconUrl", "")
    value = row.get(key, "") or ""
    return " / ".join(value) if isinstance(value, list) else value


def set_widths(ws, widths):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width


def export_raw_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.active
    ws.title = "角色字典"
    ws.freeze_panes = "A2"
    add_headers(ws, RAW_HEADERS)
    set_widths(ws, [18, 24, 18, 34, 14, 8, 12, 12, 10, 14, 20, 10, 26, 58, 58])
    for idx, row in enumerate(rows, 2):
        ws.row_dimensions[idx].height = 100
        add_avatar(ws, row, f"A{idx}")
        for col, (key, _) in enumerate(RAW_HEADERS[1:], 2):
            value = row.get(key, "")
            cell = ws.cell(row=idx, column=col, value=" / ".join(value) if isinstance(value, list) else value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def export_chinese_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("全汉化")
    ws.freeze_panes = "A2"
    add_headers(ws, CN_HEADERS)
    set_widths(ws, [18, 12, 24, 34, 14, 10, 12, 10, 12, 12, 20, 10, 26, 24, 10, 58, 58, 58, 58])
    for idx, row in enumerate(rows, 2):
        ws.row_dimensions[idx].height = 100
        add_avatar(ws, row, f"A{idx}")
        add_portrait(ws, row, f"B{idx}")
        for col, (key, _) in enumerate(CN_HEADERS[1:], 2):
            if key == "portrait":
                continue
            if key == "elementIcon":
                icon = code_icon_path(row)
                if icon:
                    img = XlsxImage(str(icon))
                    img.width = 34
                    img.height = 34
                    ws.add_image(img, f"G{idx}")
                else:
                    ws.cell(row=idx, column=col, value="缺图")
                continue
            cell = ws.cell(row=idx, column=col, value=chinese_value(row, key))
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def export_status_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("国服上线状态")
    ws.freeze_panes = "A2"
    add_headers(ws, STATUS_HEADERS)
    set_widths(ws, [18, 24, 26, 14, 34, 12, 14, 24, 34])
    for idx, row in enumerate(rows, 2):
        ws.row_dimensions[idx].height = 100
        add_avatar(ws, row, f"A{idx}")
        for col, (key, _) in enumerate(STATUS_HEADERS[1:], 2):
            value = chinese_value(row, key)
            cell = ws.cell(row=idx, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def export_asset_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("GameKee素材")
    ws.freeze_panes = "A2"
    add_headers(ws, ASSET_HEADERS)
    set_widths(ws, [18, 14, 14, 14, 24, 26, 10, 12, 10, 12, 12, 12, 12, 12, 12, 58, 58, 58])
    image_columns = {
        "teamCard": ("gamekeeLocalTeamCardFile", 70, 92, 128),
        "portrait": ("localPortraitFile", 70, 92, 128),
        "costumePortrait": ("gamekeeLocalCostumePortraitFile", 70, 92, 128),
        "gamekeeStageIcon": ("gamekeeLocalStageIconFile", 34, 38, 48),
        "gamekeeClassIcon": ("gamekeeLocalClassIconFile", 34, 38, 48),
        "gamekeeCompanyIcon": ("gamekeeLocalCompanyIconFile", 34, 38, 48),
        "gamekeeWeaponIcon": ("gamekeeLocalWeaponIconFile", 34, 38, 48),
        "gamekeeElementIcon": ("gamekeeLocalElementIconFile", 34, 38, 48),
    }
    for idx, row in enumerate(rows, 2):
        ws.row_dimensions[idx].height = 100
        add_avatar(ws, row, f"A{idx}")
        for col, (key, _) in enumerate(ASSET_HEADERS[1:], 2):
            if key in image_columns:
                path_key, width, height, size = image_columns[key]
                add_image_from_path(ws, row.get(path_key, ""), f"{chr(64 + col)}{idx}", f"{key}_{row.get('url', 'unknown')}", width, height, size)
                continue
            value = chinese_value(row, key)
            cell = ws.cell(row=idx, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def skill_text(skill: dict) -> str:
    if not skill:
        return "待确认"
    level = skill.get("level") or 1
    lines = [
        skill.get("name", "待确认"),
        f"{skill.get('type', '待确认')} / CD {skill.get('cooldown', '待确认')} / Lv{level}",
        skill.get("effect", "待确认"),
    ]
    if skill.get("treasureEffect"):
        lines += ["", "【珍藏品强化 Lv10】", skill["treasureEffect"]]
    return "\n".join(lines)


def normal_attack_text(attack: dict) -> str:
    if not attack:
        return "待确认"
    keys = [("name", "武器名"), ("type", "类型"), ("damage", "伤害"), ("effect", "普攻说明"), ("ammo", "装弹"), ("reloadTime", "换弹")]
    return "\n".join(f"{label}：{attack[key]}" for key, label in keys if attack.get(key)) or "待确认"


def _size_text(path_text: str) -> str:
    if not path_text:
        return "—"
    p = Path(path_text)
    if not p.is_file():
        return "缺文件"
    try:
        with Image.open(p) as im:
            return f"{im.width}x{im.height}"
    except Exception:  # noqa: BLE001
        return "读取失败"


def export_official_asset_sheet(wb: Workbook, rows: list[dict]) -> None:
    """官方素材（拆包资源库）挂接情况：头像 256x512 / 立绘 / HD 2048。"""
    ws = wb.create_sheet("官方素材")
    ws.freeze_panes = "A2"
    add_headers(ws, OFFICIAL_ASSET_HEADERS)
    set_widths(ws, [16, 18, 24, 16, 22, 12, 12, 12, 46, 46, 30])
    ordered = sorted(rows, key=lambda r: (not r.get("officialPortraitFile"), str(r.get("cnName") or r.get("name"))))
    idx = 2
    for row in ordered:
        portrait = row.get("officialPortraitFile", "")
        values = {
            "cnName": row.get("cnName", ""),
            "name": row.get("name", ""),
            "code": row.get("officialCode", ""),
            "status": row.get("officialAssetStatus", "未收录"),
            "portraitSize": _size_text(portrait),
            "fullbodySize": _size_text(row.get("officialFullBodyFile", "")),
            "hdSize": _size_text(row.get("officialHdFile", "")),
            "fullbodyPath": row.get("officialFullBodyFile", ""),
            "hdPath": row.get("officialHdFile", ""),
            "source": row.get("officialAssetSource", ""),
        }
        for col, (key, _) in enumerate(OFFICIAL_ASSET_HEADERS, 1):
            if key == "portrait":
                continue
            ws.cell(row=idx, column=col, value=values[key]).alignment = Alignment(vertical="center", wrap_text=False)
        ws.row_dimensions[idx].height = 52
        add_image_from_path(ws, portrait, f"A{idx}", f"official_{row.get('url', idx)}", 24, 48, 64)
        idx += 1


def export_skill_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("技能组")
    ws.freeze_panes = "A2"
    add_headers(ws, SKILL_HEADERS)
    set_widths(ws, [18, 24, 18, 12, 12, 30, 48, 48, 48, 50])
    for idx, row in enumerate(rows, 2):
        ws.row_dimensions[idx].height = 130
        values = {
            "cnName": row.get("cnName", ""),
            "name": row.get("name", ""),
            "gamekeeSquadCn": row.get("gamekeeSquadCn", "待确认"),
            "gamekeeStageCn": row.get("gamekeeStageCn", "待确认"),
            "gamekeeWeaponCn": row.get("gamekeeWeaponCn", "待确认"),
            "normalAttack": normal_attack_text(row.get("normalAttack", {})),
            "skill1": skill_text(row.get("skills", {}).get("skill1", {})),
            "skill2": skill_text(row.get("skills", {}).get("skill2", {})),
            "burstSkill": skill_text(row.get("skills", {}).get("burstSkill", {})),
            "sourceUrl": row.get("source", {}).get("url", ""),
        }
        for col, (key, _) in enumerate(SKILL_HEADERS, 1):
            cell = ws.cell(row=idx, column=col, value=values[key])
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_favorite_item_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("珍藏品")
    ws.freeze_panes = "A2"
    add_headers(ws, FAVORITE_ITEM_HEADERS)
    set_widths(ws, [20, 24, 14, 30, 14, 52, 72, 58])
    selected = [row for row in rows if row.get("favoriteItem", {}).get("status") == "已确认"]
    for idx, row in enumerate(selected, 2):
        item = row["favoriteItem"]
        favorite_skills = item.get("skills", {})
        skill_summary = "\n\n".join(
            f"{skill.get('name', key)}：{skill.get('effect', '待确认')}"
            for key, skill in favorite_skills.items()
        ) or "待确认"
        values = {"cnName": row.get("cnName", ""), "name": row.get("name", ""), **item, "skills": skill_summary}
        for col, (key, _) in enumerate(FAVORITE_ITEM_HEADERS, 1):
            cell = ws.cell(row=idx, column=col, value=values.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_costume_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("\u76ae\u80a4\u6e05\u5355")
    ws.freeze_panes = "A2"
    add_headers(ws, COSTUME_HEADERS)
    set_widths(ws, [20, 24, 18, 30, 14, 16, 58, 58])
    index = 2
    for row in rows:
        for costume in row.get("costumes", []):
            ws.row_dimensions[index].height = 106
            values = {"parentCnName": row.get("cnName", ""), "parentName": row.get("name", ""), **costume}
            add_image_from_path(ws, costume.get("localPortraitFile", ""), f"C{index}", f"costume_{row.get('url', 'unknown')}_{costume.get('contentId', '')}", 44, 96, 128)
            for col, (key, _) in enumerate(COSTUME_HEADERS, 1):
                if key == "portrait":
                    continue
                cell = ws.cell(row=index, column=col, value=values.get(key, ""))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            index += 1


def export_non_playable_sheet(wb: Workbook) -> None:
    if not NON_PLAYABLE_CARDS.exists():
        return
    cards = json.loads(NON_PLAYABLE_CARDS.read_text(encoding="utf-8"))
    ws = wb.create_sheet("NPC与剧情角色")
    ws.freeze_panes = "A2"
    add_headers(ws, NON_PLAYABLE_HEADERS)
    set_widths(ws, [18, 22, 22, 14, 18, 12, 12, 12, 12, 66, 58])
    for index, card in enumerate(cards, 2):
        ws.row_dimensions[index].height = 116
        add_image_from_path(ws, card.get("localPortraitFile", ""), f"A{index}", f"non_playable_{card.get('gamekeeContentId', '')}", 54, 100, 128)
        values = {**card, "aliases": " / ".join(card.get("aliases", [])), "skills": "\n\n".join(skill_text(skill) for skill in card.get("skills", {}).values()), "sourceUrl": card.get("source", {}).get("url", "")}
        for col, (key, _) in enumerate(NON_PLAYABLE_HEADERS[1:], 2):
            cell = ws.cell(row=index, column=col, value=values.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_xlsx(rows: list[dict]) -> Path:
    wb = Workbook()
    export_raw_sheet(wb, rows)
    export_chinese_sheet(wb, rows)
    export_status_sheet(wb, rows)
    export_asset_sheet(wb, rows)
    export_costume_sheet(wb, rows)
    export_non_playable_sheet(wb)
    export_skill_sheet(wb, rows)
    export_favorite_item_sheet(wb, rows)
    export_official_asset_sheet(wb, rows)
    for ws in wb.worksheets:
        finish_sheet(ws)
    try:
        wb.save(XLSX_OUT)
        return XLSX_OUT
    except PermissionError:
        wb.save(XLSX_FALLBACK_OUT)
        return XLSX_FALLBACK_OUT


def main() -> None:
    global ENHANCED_DICTIONARY, XLSX_OUT, HTML_OUT
    parser = argparse.ArgumentParser(description="Export NIKKE character dictionary tables")
    parser.add_argument("--input", type=Path, default=ENHANCED_DICTIONARY)
    parser.add_argument("--xlsx-output", type=Path, default=XLSX_OUT)
    parser.add_argument("--html-output", type=Path, default=HTML_OUT)
    args = parser.parse_args()
    ENHANCED_DICTIONARY = args.input
    XLSX_OUT = args.xlsx_output
    HTML_OUT = args.html_output
    payload = json.loads(ENHANCED_DICTIONARY.read_text(encoding="utf-8"))
    rows = payload.get("characters", payload) if isinstance(payload, dict) else payload
    rows = normalize_rows(rows)
    export_html(rows)
    xlsx_out = export_xlsx(rows)
    print(f"HTML: {HTML_OUT}")
    print(f"XLSX: {xlsx_out}")


if __name__ == "__main__":
    main()
